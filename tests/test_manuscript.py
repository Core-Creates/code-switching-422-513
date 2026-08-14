"""The manuscript is generated from the artifacts, so its numbers are testable."""
import os
import re

import docx
import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOC = os.path.join(ROOT, "paper", "code_switching_paper_v2.docx")


@pytest.fixture(scope="module")
def text():
    if not os.path.exists(DOC):
        pytest.skip("manuscript not generated yet; run paper/generate_manuscript.py")
    d = docx.Document(DOC)
    body = "\n".join(p.text for p in d.paragraphs)
    for t in d.tables:
        for row in t.rows:
            body += "\n" + " | ".join(c.text for c in row.cells)
    return body


def test_all_three_authors_on_the_byline(text):
    for name in ("Corrina Alcoser", "Keeban Villarreal", "Michael Pendleton"):
        assert name in text


def test_key_numbers_match_the_artifacts(text):
    import json
    e2e = json.load(open(os.path.join(ROOT, "results", "end_to_end.json")))
    num = json.load(open(os.path.join(ROOT, "results", "numerics.json")))
    total = e2e["X"]["mechanisms"] + e2e["Z"]["mechanisms"]
    assert f"{total:,}" in text
    assert f"{num['certified']['slope']:.2f}" in text
    assert f"{num['M1 flag removed']['slope']:.2f}" in text
    assert "97,466" in text


def test_banned_vocabulary_absent(text):
    """No estimate stands in for a number we possess."""
    for banned in ("approximately 15", "~15", "about 15 CNOT"):
        assert banned not in text


def test_no_claim_of_the_disproved_theorem(text):
    """The weight-one formulation is false; it must not reappear as a claim."""
    assert "at most a weight-one error on the output" not in text
    assert "unattainable" in text or "cannot be certified" in text


def test_limitations_are_stated(text):
    assert "Limitations" in text
    for phrase in ("capped", "single faults"):
        assert phrase in text
